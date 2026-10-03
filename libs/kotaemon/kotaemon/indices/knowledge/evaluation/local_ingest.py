"""Prepare deterministic, local-only parsing and chunk review drafts."""

from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from dataclasses import asdict, dataclass, replace
from pathlib import Path
from typing import Any, Sequence

from .local_corpus import (
    SourceFile,
    canonical_sources,
    scan_sources,
    unsupported_paths,
)

_CHUNK_CONFIG_VERSION = "main-token-only-v1"
_CHUNK_SIZE = 1024
_CHUNK_OVERLAP = 256
_CHUNK_SEPARATOR = "\n\n"
_CHUNK_BACKUP_SEPARATORS = ("\n", ".", " ", "\u200b")
_LOCATOR_KEYS = frozenset(
    {
        "page_label",
        "page_number",
        "sheet_name",
        "row_number",
        "category",
        "category_depth",
    }
)


@dataclass(frozen=True)
class SourceUnit:
    """One reader-returned unit with root-relative provenance and stable locator."""

    unit_id: str
    source_id: str
    relative_path: str
    unit_ordinal: int
    locator: dict[str, str | int | float | bool]
    text: str
    normalized_text: str


@dataclass(frozen=True)
class DraftChunk:
    """A baseline token chunk with offsets into its normalized source unit."""

    chunk_id: str
    source_id: str
    relative_path: str
    unit_id: str
    unit_ordinal: int
    chunk_ordinal: int
    locator: dict[str, str | int | float | bool]
    text: str
    char_start: int
    char_end: int


@dataclass(frozen=True)
class TopicCandidate:
    """A content-derived heading suggestion that still needs human review."""

    candidate_id: str
    source_id: str
    relative_path: str
    unit_id: str
    text: str
    locator: dict[str, str | int | float | bool]
    needs_review: bool = True


@dataclass(frozen=True)
class ExcludedInput:
    relative_path: str
    source_id: str | None
    reason: str
    duplicate_of: str | None = None


@dataclass(frozen=True)
class ExtractionFailure:
    relative_path: str
    source_id: str
    error_type: str
    reason: str = "reader_failed"


@dataclass(frozen=True)
class ReaderDiagnostic:
    relative_path: str
    source_id: str
    reader_used: str | None
    attempted_readers: tuple[str, ...]
    fallback_used: bool
    granularity: str
    needs_review: bool = False
    detail: str | None = None


@dataclass(frozen=True)
class QualityLocator:
    relative_path: str
    source_id: str
    locator: dict[str, str | int | float | bool]
    reason: str


@dataclass(frozen=True)
class PageQuality:
    relative_path: str
    source_id: str
    total_pages: int
    extracted_pages: int
    empty_pages: int


@dataclass(frozen=True)
class SheetQuality:
    relative_path: str
    source_id: str
    sheet_name: str
    extracted_rows: int
    blank_rows: int


@dataclass(frozen=True)
class ChunkMappingIssue:
    relative_path: str
    source_id: str
    unit_id: str
    locator: dict[str, str | int | float | bool]
    reason: str
    error_type: str
    needs_review: bool = True


@dataclass(frozen=True)
class ExtractionQuality:
    empty_locators: tuple[QualityLocator, ...]
    unusable_locators: tuple[QualityLocator, ...]
    reader_diagnostics: tuple[ReaderDiagnostic, ...]
    extraction_failures: tuple[ExtractionFailure, ...]
    chunk_mapping_issues: tuple[ChunkMappingIssue, ...]
    pages: tuple[PageQuality, ...]
    sheets: tuple[SheetQuality, ...]


@dataclass(frozen=True)
class DraftSummary:
    parsed_source_ids: tuple[str, ...]
    topic_candidates: tuple[TopicCandidate, ...]
    topic_review_source_ids: tuple[str, ...]
    quality: ExtractionQuality
    source_units: tuple[SourceUnit, ...]
    chunks: tuple[DraftChunk, ...]
    excluded_inputs: tuple[ExcludedInput, ...]
    records_payload: bytes
    anchors_payload: bytes

    @property
    def chunk_ids(self) -> tuple[str, ...]:
        return tuple(chunk.chunk_id for chunk in self.chunks)


@dataclass(frozen=True)
class _SplitPart:
    text: str
    start: int
    end: int
    token_count: int


def _canonical_json(value: Any) -> bytes:
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _normalize_text(text: str) -> str:
    line_normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    return unicodedata.normalize("NFC", line_normalized)


def _reader_diagnostic(
    source: SourceFile,
    reader_used: str | None,
    attempted_readers: tuple[str, ...],
    fallback_used: bool,
    granularity: str,
    *,
    needs_review: bool = False,
    detail: str | None = None,
) -> ReaderDiagnostic:
    return ReaderDiagnostic(
        relative_path=source.relative_path,
        source_id=source.source_id,
        reader_used=reader_used,
        attempted_readers=attempted_readers,
        fallback_used=fallback_used,
        granularity=granularity,
        needs_review=needs_review,
        detail=detail,
    )


def _load_reader_documents(
    source: SourceFile,
    path: Path,
) -> tuple[list[Any], ReaderDiagnostic, str | None]:
    """Load locally and return a diagnostic that names the parser actually used."""
    if source.suffix == ".pdf":
        reader_name = "PDFReader"
        diagnostic = _reader_diagnostic(
            source, reader_name, (reader_name,), False, "page"
        )
        try:
            from llama_index.readers.file import PDFReader

            return PDFReader().load_data(file=path), diagnostic, None
        except Exception as exc:
            return (
                [],
                replace(diagnostic, needs_review=True, detail="reader failed"),
                type(exc).__name__,
            )

    if source.suffix == ".docx":
        unstructured_name = "UnstructuredReader(split_documents=True)"
        fallback_name = "DocxReader(local fallback)"
        try:
            from kotaemon.loaders.unstructured_loader import UnstructuredReader

            documents = UnstructuredReader().load_data(file=path, split_documents=True)
            diagnostic = _reader_diagnostic(
                source, unstructured_name, (unstructured_name,), False, "element"
            )
            return documents, diagnostic, None
        except ImportError:
            # Unstructured's local DOCX partitioning requires system libmagic.
            # Use Kotaemon's existing python-docx reader when that optional
            # system library is absent.
            try:
                from kotaemon.loaders.docx_loader import DocxReader

                documents = DocxReader().load_data(file_path=path)
                diagnostic = _reader_diagnostic(
                    source,
                    fallback_name,
                    (unstructured_name, fallback_name),
                    True,
                    "document",
                    needs_review=True,
                    detail=(
                        "Local Unstructured partitioning raised ImportError; "
                        "DocxReader coalesces paragraph text into one document."
                    ),
                )
                return documents, diagnostic, None
            except Exception as exc:
                diagnostic = _reader_diagnostic(
                    source,
                    None,
                    (unstructured_name, fallback_name),
                    True,
                    "document",
                    needs_review=True,
                    detail="Unstructured and DocxReader fallback both failed",
                )
                return [], diagnostic, type(exc).__name__
        except Exception as exc:
            diagnostic = _reader_diagnostic(
                source,
                None,
                (unstructured_name,),
                False,
                "element",
                needs_review=True,
                detail="reader failed",
            )
            return [], diagnostic, type(exc).__name__

    if source.suffix == ".xlsx":
        reader_name = "ExcelRowReader"
        diagnostic = _reader_diagnostic(
            source, reader_name, (reader_name,), False, "row"
        )
        try:
            from kotaemon.loaders.excel_loader import ExcelRowReader

            return ExcelRowReader().load_data(file=path), diagnostic, None
        except Exception as exc:
            return (
                [],
                replace(diagnostic, needs_review=True, detail="reader failed"),
                type(exc).__name__,
            )

    if source.suffix == ".md":
        reader_name = "TxtReader"
        diagnostic = _reader_diagnostic(
            source, reader_name, (reader_name,), False, "file"
        )
        try:
            from kotaemon.loaders.txt_loader import TxtReader

            return TxtReader().load_data(file_path=path), diagnostic, None
        except Exception as exc:
            return (
                [],
                replace(diagnostic, needs_review=True, detail="reader failed"),
                type(exc).__name__,
            )

    raise ValueError(f"Unsupported local source suffix: {source.suffix}")


def _locator_for(
    metadata: dict[str, Any], suffix: str, ordinal: int, reader_used: str | None
):
    locator: dict[str, str | int | float | bool] = {}
    for key in _LOCATOR_KEYS:
        value = metadata.get(key)
        if isinstance(value, (str, int, float, bool)):
            locator[key] = value
    if suffix == ".docx" and reader_used == "UnstructuredReader(split_documents=True)":
        locator["element_ordinal"] = ordinal
    else:
        locator["unit_ordinal"] = ordinal
    return locator


def _is_unusable_locator(locator: dict[str, Any], suffix: str) -> bool:
    if suffix == ".pdf":
        return not any(key in locator for key in ("page_label", "page_number"))
    if suffix == ".docx":
        return not any(
            key in locator
            for key in ("category", "category_depth", "page_number", "page_label")
        )
    if suffix == ".xlsx":
        return not {"sheet_name", "row_number"}.issubset(locator)
    return False


def _source_unit(
    source: SourceFile,
    document: Any,
    ordinal: int,
    reader_used: str | None,
) -> SourceUnit:
    text = str(getattr(document, "text", "") or "")
    metadata = getattr(document, "metadata", None) or {}
    locator = _locator_for(metadata, source.suffix, ordinal, reader_used)
    unit_id = _digest(
        b"local-unit:v1\0"
        + source.sha256.encode("ascii")
        + b"\0"
        + str(ordinal).encode("ascii")
        + b"\0"
        + _canonical_json(locator)
    )
    return SourceUnit(
        unit_id=unit_id,
        source_id=source.source_id,
        relative_path=source.relative_path,
        unit_ordinal=ordinal,
        locator=locator,
        text=text,
        normalized_text=_normalize_text(text),
    )


def _split_with_offsets(text: str, splitter: Any) -> list[tuple[str, int, int]]:
    """Mirror TokenTextSplitter merging while retaining exact source offsets."""
    if not text:
        return []

    token_splitter = splitter._obj
    raw_splits = token_splitter._split(text, chunk_size=_CHUNK_SIZE)
    parts: list[_SplitPart] = []
    cursor = 0
    for raw_text in raw_splits:
        start = text.find(raw_text, cursor)
        if start < 0:
            raise ValueError(
                "Token splitter returned a segment outside its source unit"
            )
        end = start + len(raw_text)
        parts.append(
            _SplitPart(
                text=raw_text,
                start=start,
                end=end,
                token_count=len(token_splitter._tokenizer(raw_text)),
            )
        )
        cursor = end
    if cursor != len(text):
        raise ValueError("Token splitter segments did not cover their source unit")

    chunks: list[tuple[str, int, int]] = []
    current: list[_SplitPart] = []
    current_tokens = 0

    def append_current() -> None:
        joined = "".join(part.text for part in current)
        chunk_text = joined.strip()
        if not chunk_text:
            return
        left_trim = len(joined) - len(joined.lstrip())
        right_edge = len(joined.rstrip())
        start = current[0].start + left_trim
        right_trim = len(joined) - right_edge
        end = current[-1].end - right_trim
        if not (0 <= start <= end <= len(text)) or text[start:end] != chunk_text:
            raise ValueError("Token splitter chunk could not be mapped exactly")
        chunks.append((chunk_text, start, end))

    for part in parts:
        if current_tokens + part.token_count > _CHUNK_SIZE:
            append_current()
            while current and (
                current_tokens > _CHUNK_OVERLAP
                or current_tokens + part.token_count > _CHUNK_SIZE
            ):
                first = current.pop(0)
                current_tokens -= first.token_count
        current.append(part)
        current_tokens += part.token_count

    append_current()
    return chunks


def _token_splitter():
    from kotaemon.indices.splitters import TokenSplitter

    return TokenSplitter(
        chunk_size=_CHUNK_SIZE,
        chunk_overlap=_CHUNK_OVERLAP,
        separator=_CHUNK_SEPARATOR,
        backup_separators=list(_CHUNK_BACKUP_SEPARATORS),
    )


def _topic_candidates(unit: SourceUnit, suffix: str) -> tuple[TopicCandidate, ...]:
    candidates: list[TopicCandidate] = []
    if suffix == ".md":
        for line_number, line in enumerate(unit.text.splitlines(), 1):
            match = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", line)
            if not match:
                continue
            title = match.group(1).strip()
            if not title:
                continue
            locator = {**unit.locator, "heading_line": line_number}
            candidate_id = _digest(
                b"local-topic:v1\0"
                + unit.unit_id.encode("ascii")
                + b"\0"
                + str(line_number).encode("ascii")
                + b"\0"
                + title.encode("utf-8")
            )
            candidates.append(
                TopicCandidate(
                    candidate_id=candidate_id,
                    source_id=unit.source_id,
                    relative_path=unit.relative_path,
                    unit_id=unit.unit_id,
                    text=title,
                    locator=locator,
                )
            )
    elif isinstance(unit.locator.get("category"), str) and unit.locator[
        "category"
    ].lower() in {"title", "header"}:
        title = unit.normalized_text
        if title:
            candidate_id = _digest(
                b"local-topic:v1\0"
                + unit.unit_id.encode("ascii")
                + b"\0"
                + title.encode("utf-8")
            )
            candidates.append(
                TopicCandidate(
                    candidate_id=candidate_id,
                    source_id=unit.source_id,
                    relative_path=unit.relative_path,
                    unit_id=unit.unit_id,
                    text=title,
                    locator=unit.locator,
                )
            )
    return tuple(candidates)


def _xlsx_quality(
    source: SourceFile,
    path: Path,
    documents: list[Any],
) -> tuple[SheetQuality, ...]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        quality = []
        extracted_by_sheet: dict[str, int] = {}
        for document in documents:
            metadata = getattr(document, "metadata", None) or {}
            sheet_name = metadata.get("sheet_name")
            if isinstance(sheet_name, str):
                extracted_by_sheet[sheet_name] = (
                    extracted_by_sheet.get(sheet_name, 0) + 1
                )

        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(min_row=2, values_only=True))
            blank_rows = sum(
                1
                for row in rows
                if not any(value is not None and str(value).strip() for value in row)
            )
            quality.append(
                SheetQuality(
                    relative_path=source.relative_path,
                    source_id=source.source_id,
                    sheet_name=sheet.title,
                    extracted_rows=extracted_by_sheet.get(sheet.title, 0),
                    blank_rows=blank_rows,
                )
            )
        return tuple(quality)
    finally:
        workbook.close()


def _make_chunk_id(
    unit: SourceUnit,
    chunk_ordinal: int,
    chunk_text: str,
) -> str:
    return _digest(
        b"local-chunk:v1\0"
        + _canonical_json(
            {
                "version": _CHUNK_CONFIG_VERSION,
                "name": "TokenSplitter",
                "chunk_size": _CHUNK_SIZE,
                "chunk_overlap": _CHUNK_OVERLAP,
                "separator": _CHUNK_SEPARATOR,
                "backup_separators": list(_CHUNK_BACKUP_SEPARATORS),
            }
        )
        + b"\0"
        + unit.source_id.encode("ascii")
        + b"\0"
        + _canonical_json(unit.locator)
        + b"\0"
        + str(unit.unit_ordinal).encode("ascii")
        + b"\0"
        + str(chunk_ordinal).encode("ascii")
        + b"\0"
        + _digest(chunk_text.encode("utf-8")).encode("ascii")
    )


def _serialize_quality(quality: ExtractionQuality) -> dict[str, Any]:
    return asdict(quality)


def build_local_draft(
    root: Path,
    output_dir: Path,
    *,
    sample_ids: Sequence[str] | None = None,
) -> DraftSummary:
    """Parse unique supported files locally and write deterministic review drafts.

    ``sample_ids`` selects source SHA-256 IDs from ``scan_sources``. Duplicate
    byte paths remain in provenance, while only their canonical path is parsed.
    The function writes its two payloads only inside the caller-provided output
    directory and never creates query-linked anchors.
    """
    root = Path(root)
    output_dir = Path(output_dir)
    inventory = scan_sources(root)
    canonical = canonical_sources(inventory)
    all_ids = {source.source_id for source in canonical}
    selected_ids = all_ids if sample_ids is None else set(sample_ids)
    unknown_ids = selected_ids - all_ids
    if unknown_ids:
        raise ValueError(f"Unknown sample source IDs: {', '.join(sorted(unknown_ids))}")
    selected = tuple(source for source in canonical if source.source_id in selected_ids)

    excluded_inputs = [
        ExcludedInput(
            relative_path=source.relative_path,
            source_id=source.source_id,
            duplicate_of=source.duplicate_of,
            reason=(
                "duplicate_bytes"
                if source.duplicate_of is not None and source.source_id in selected_ids
                else (
                    "not_selected"
                    if source.source_id not in selected_ids
                    else "duplicate_bytes"
                )
            ),
        )
        for source in inventory
        if source.duplicate_of is not None or source.source_id not in selected_ids
    ]
    excluded_inputs.extend(
        ExcludedInput(
            relative_path=relative_path,
            source_id=None,
            reason="unsupported_suffix",
        )
        for relative_path in unsupported_paths(root)
    )
    excluded_inputs.sort(key=lambda item: item.relative_path)

    source_units: list[SourceUnit] = []
    chunks: list[DraftChunk] = []
    topic_candidates: list[TopicCandidate] = []
    empty_locators: list[QualityLocator] = []
    unusable_locators: list[QualityLocator] = []
    extraction_failures: list[ExtractionFailure] = []
    chunk_mapping_issues: list[ChunkMappingIssue] = []
    page_quality: list[PageQuality] = []
    sheet_quality: list[SheetQuality] = []
    parsed_source_ids: list[str] = []
    failed_source_ids: set[str] = set()
    reader_diagnostic_by_id: dict[str, ReaderDiagnostic] = {}
    reader_diagnostics: list[ReaderDiagnostic] = []
    splitter = None

    for source in selected:
        path = root / source.relative_path
        documents, reader_diagnostic, error_type = _load_reader_documents(source, path)
        reader_diagnostics.append(reader_diagnostic)
        reader_diagnostic_by_id[source.source_id] = reader_diagnostic
        if error_type is not None:
            failed_source_ids.add(source.source_id)
            extraction_failures.append(
                ExtractionFailure(
                    relative_path=source.relative_path,
                    source_id=source.source_id,
                    error_type=error_type,
                )
            )
            continue

        if not documents:
            failed_source_ids.add(source.source_id)
            reader_diagnostics[-1] = replace(
                reader_diagnostic,
                needs_review=True,
                detail="reader returned no source units",
            )
            reader_diagnostic_by_id[source.source_id] = reader_diagnostics[-1]
            extraction_failures.append(
                ExtractionFailure(
                    relative_path=source.relative_path,
                    source_id=source.source_id,
                    error_type="EmptyReaderResult",
                    reason="no_source_units_returned",
                )
            )
            continue

        parsed_source_ids.append(source.source_id)
        source_units_for_file: list[SourceUnit] = []
        for ordinal, document in enumerate(documents):
            unit = _source_unit(
                source,
                document,
                ordinal,
                reader_diagnostic.reader_used,
            )
            source_units.append(unit)
            source_units_for_file.append(unit)

            if not unit.normalized_text.strip():
                empty_locators.append(
                    QualityLocator(
                        relative_path=unit.relative_path,
                        source_id=unit.source_id,
                        locator=unit.locator,
                        reason="empty_extraction",
                    )
                )
            if _is_unusable_locator(unit.locator, source.suffix):
                unusable_locators.append(
                    QualityLocator(
                        relative_path=unit.relative_path,
                        source_id=unit.source_id,
                        locator=unit.locator,
                        reason="reader_locator_unavailable",
                    )
                )
            topic_candidates.extend(_topic_candidates(unit, source.suffix))

            if unit.normalized_text.strip():
                try:
                    if splitter is None:
                        splitter = _token_splitter()
                    unit_chunks = _split_with_offsets(unit.normalized_text, splitter)
                except Exception as exc:
                    chunk_mapping_issues.append(
                        ChunkMappingIssue(
                            relative_path=unit.relative_path,
                            source_id=unit.source_id,
                            unit_id=unit.unit_id,
                            locator=unit.locator,
                            reason="offset_mapping_failed",
                            error_type=type(exc).__name__,
                        )
                    )
                    continue

                for chunk_ordinal, (text, start, end) in enumerate(unit_chunks):
                    chunks.append(
                        DraftChunk(
                            chunk_id=_make_chunk_id(unit, chunk_ordinal, text),
                            source_id=unit.source_id,
                            relative_path=unit.relative_path,
                            unit_id=unit.unit_id,
                            unit_ordinal=unit.unit_ordinal,
                            chunk_ordinal=chunk_ordinal,
                            locator=unit.locator,
                            text=text,
                            char_start=start,
                            char_end=end,
                        )
                    )

        if source.suffix == ".pdf":
            empty_pages = sum(
                not unit.normalized_text.strip() for unit in source_units_for_file
            )
            page_quality.append(
                PageQuality(
                    relative_path=source.relative_path,
                    source_id=source.source_id,
                    total_pages=len(source_units_for_file),
                    extracted_pages=len(source_units_for_file) - empty_pages,
                    empty_pages=empty_pages,
                )
            )
        elif source.suffix == ".xlsx":
            try:
                sheet_quality.extend(_xlsx_quality(source, path, documents))
            except Exception as exc:
                extraction_failures.append(
                    ExtractionFailure(
                        relative_path=source.relative_path,
                        source_id=source.source_id,
                        error_type=type(exc).__name__,
                        reason="quality_count_failed",
                    )
                )

    topic_source_ids = {candidate.source_id for candidate in topic_candidates}
    topic_review_source_ids = tuple(
        source.source_id
        for source in selected
        if source.source_id not in topic_source_ids
    )
    quality = ExtractionQuality(
        empty_locators=tuple(empty_locators),
        unusable_locators=tuple(unusable_locators),
        reader_diagnostics=tuple(reader_diagnostics),
        extraction_failures=tuple(extraction_failures),
        chunk_mapping_issues=tuple(chunk_mapping_issues),
        pages=tuple(page_quality),
        sheets=tuple(sheet_quality),
    )

    status_by_id = {
        source.source_id: (
            "failed" if source.source_id in failed_source_ids else "parsed"
        )
        for source in selected
    }
    excluded_inputs = tuple(excluded_inputs)
    source_records = [
        {
            "source_id": source.source_id,
            "relative_path": source.relative_path,
            "sha256": source.sha256,
            "suffix": source.suffix,
            "byte_size": source.byte_size,
            "duplicate_of": source.duplicate_of,
            "selected_reader": (
                reader_diagnostic_by_id[source.source_id].reader_used
                if source.duplicate_of is None
                and source.source_id in reader_diagnostic_by_id
                else None
            ),
            "reader_attempts": (
                list(reader_diagnostic_by_id[source.source_id].attempted_readers)
                if source.duplicate_of is None
                and source.source_id in reader_diagnostic_by_id
                else []
            ),
            "reader_fallback_used": (
                reader_diagnostic_by_id[source.source_id].fallback_used
                if source.duplicate_of is None
                and source.source_id in reader_diagnostic_by_id
                else False
            ),
            "reader_granularity": (
                reader_diagnostic_by_id[source.source_id].granularity
                if source.duplicate_of is None
                and source.source_id in reader_diagnostic_by_id
                else None
            ),
            "status": (
                "duplicate_bytes"
                if source.duplicate_of is not None
                else status_by_id.get(source.source_id, "not_selected")
            ),
        }
        for source in inventory
    ]
    records = {
        "schema_version": 1,
        "parser_configuration": {
            ".pdf": "llama_index.readers.file.PDFReader",
            ".docx": (
                "UnstructuredReader(split_documents=True, local); "
                "DocxReader fallback after local Unstructured ImportError"
            ),
            ".md": "TxtReader",
            ".xlsx": "ExcelRowReader",
        },
        "splitter_configuration": {
            "version": _CHUNK_CONFIG_VERSION,
            "name": "TokenSplitter",
            "chunk_size": _CHUNK_SIZE,
            "chunk_overlap": _CHUNK_OVERLAP,
            "separator": _CHUNK_SEPARATOR,
            "backup_separators": list(_CHUNK_BACKUP_SEPARATORS),
        },
        "sources": source_records,
        "source_units": [asdict(unit) for unit in source_units],
        "chunks": [asdict(chunk) for chunk in chunks],
        "topic_candidates": [asdict(candidate) for candidate in topic_candidates],
        "topic_review_source_ids": list(topic_review_source_ids),
        "quality": _serialize_quality(quality),
        "excluded_inputs": [asdict(item) for item in excluded_inputs],
    }
    records_payload = _canonical_json(records) + b"\n"
    anchors_payload = b""

    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "records.json").write_bytes(records_payload)
    (output_dir / "anchors.jsonl").write_bytes(anchors_payload)

    return DraftSummary(
        parsed_source_ids=tuple(parsed_source_ids),
        topic_candidates=tuple(topic_candidates),
        topic_review_source_ids=topic_review_source_ids,
        quality=quality,
        source_units=tuple(source_units),
        chunks=tuple(chunks),
        excluded_inputs=excluded_inputs,
        records_payload=records_payload,
        anchors_payload=anchors_payload,
    )
