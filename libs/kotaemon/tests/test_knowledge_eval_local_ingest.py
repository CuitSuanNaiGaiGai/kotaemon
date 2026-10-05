"""Generated-file tests for local parsing and chunk draft preparation."""

from __future__ import annotations

import json
import shutil
import unicodedata
from pathlib import Path

import fitz
from docx import Document as WordDocument
from openpyxl import Workbook

from kotaemon.indices.knowledge.evaluation.local_ingest import build_local_draft


def _make_multiformat_corpus(root: Path) -> Path:
    """Create local synthetic inputs without touching the retained source corpus."""
    root.mkdir(parents=True, exist_ok=True)

    markdown = root / "research" / "project-notes.md"
    markdown.parent.mkdir(parents=True)
    markdown.write_text(
        "# Project Notes\n\n" + "evidence token " * 1200,
        encoding="utf-8",
    )
    shutil.copyfile(markdown, root / "zz-duplicate.md")
    (root / "empty.md").write_text("", encoding="utf-8")
    (root / "ignored.py").write_text("unsupported generated input", encoding="utf-8")
    (root / "~$draft.xlsx").write_bytes(b"generated office lock")
    (root / ".DS_Store").write_bytes(b"generated metadata file")
    (root / "ignored-link.py").symlink_to(root / "ignored.py")

    pdf_path = root / "report.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "PDF evidence on the first page.")
    pdf.new_page()
    pdf.save(pdf_path)
    pdf.close()

    docx_path = root / "summary.docx"
    word = WordDocument()
    word.add_heading("Quarterly Summary", level=1)
    word.add_paragraph("DOCX evidence remains associated with its reader element.")
    word.save(docx_path)
    (root / "broken.docx").write_bytes(b"not a valid generated DOCX package")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Roster"
    sheet.append(["Person", "Evidence"])
    sheet.append(["Ada", "First spreadsheet row"])
    sheet.append([None, None])
    sheet.append(["Grace", "Second spreadsheet row"])
    workbook.save(root / "people.xlsx")
    workbook.close()
    return root


def test_build_local_draft_is_stable_and_reports_empty_or_failed_extraction(
    tmp_path, monkeypatch
):
    root = _make_multiformat_corpus(tmp_path / "sources")
    unsupported = root / "ignored.py"
    original_open = Path.open

    def reject_unsupported_open(path, *args, **kwargs):
        if path == unsupported:
            raise AssertionError("unsupported input must not be opened")
        return original_open(path, *args, **kwargs)

    monkeypatch.setattr(Path, "open", reject_unsupported_open)

    first = build_local_draft(root, tmp_path / "draft-a")
    second = build_local_draft(root, tmp_path / "draft-b")

    assert first.records_payload == second.records_payload
    assert first.anchors_payload == second.anchors_payload == b""
    assert (tmp_path / "draft-a" / "records.json").read_bytes() == first.records_payload
    assert (tmp_path / "draft-a" / "anchors.jsonl").read_bytes() == b""
    assert not (tmp_path / "draft-a" / "anchors.json").exists()
    assert first.quality.empty_locators
    assert any(
        item.relative_path == "report.pdf" and item.locator.get("page_label") == "2"
        for item in first.quality.empty_locators
    )
    assert any(
        sheet.relative_path == "people.xlsx"
        and sheet.sheet_name == "Roster"
        and sheet.extracted_rows == 2
        and sheet.blank_rows == 1
        for sheet in first.quality.sheets
    )
    assert any(
        failure.relative_path == "broken.docx"
        for failure in first.quality.extraction_failures
    )
    exclusions = {item.relative_path: item for item in first.excluded_inputs}
    assert exclusions["ignored.py"].source_id is None
    assert exclusions["ignored.py"].reason == "unsupported_suffix"
    assert "~$draft.xlsx" not in exclusions
    assert ".DS_Store" not in exclusions
    assert "ignored-link.py" not in exclusions


def test_build_local_draft_preserves_locators_offsets_and_unique_byte_parsing(tmp_path):
    root = _make_multiformat_corpus(tmp_path / "sources")
    first = build_local_draft(root, tmp_path / "draft-a")
    second = build_local_draft(root, tmp_path / "draft-b")

    assert first.parsed_source_ids == second.parsed_source_ids
    assert [chunk.chunk_id for chunk in first.chunks] == [
        chunk.chunk_id for chunk in second.chunks
    ]
    assert len(first.parsed_source_ids) == len(set(first.parsed_source_ids))
    assert any(
        item.relative_path == "zz-duplicate.md" and item.reason == "duplicate_bytes"
        for item in first.excluded_inputs
    )
    assert any(
        candidate.text == "Project Notes"
        and candidate.relative_path == "research/project-notes.md"
        for candidate in first.topic_candidates
    )
    assert str(root) not in first.records_payload.decode("utf-8")
    record = json.loads(first.records_payload)
    source_records = {source["relative_path"]: source for source in record["sources"]}

    canonical_markdown_id = source_records["research/project-notes.md"]["source_id"]
    duplicate_parse_diagnostics = [
        diagnostic
        for diagnostic in first.quality.reader_diagnostics
        if diagnostic.source_id == canonical_markdown_id
    ]
    assert len(duplicate_parse_diagnostics) == 1
    assert duplicate_parse_diagnostics[0].relative_path == "research/project-notes.md"
    assert "zz-duplicate.md" not in {
        diagnostic.relative_path for diagnostic in first.quality.reader_diagnostics
    }

    docx_diagnostic = next(
        diagnostic
        for diagnostic in first.quality.reader_diagnostics
        if diagnostic.relative_path == "summary.docx"
    )
    assert (
        source_records["summary.docx"]["selected_reader"] == docx_diagnostic.reader_used
    )
    assert source_records["summary.docx"]["reader_attempts"] == list(
        docx_diagnostic.attempted_readers
    )
    if docx_diagnostic.fallback_used:
        assert docx_diagnostic.granularity == "document"
        assert docx_diagnostic.needs_review
        assert "coalesces paragraph text" in docx_diagnostic.detail
    else:
        assert docx_diagnostic.granularity == "element"
    empty_source_id = source_records["empty.md"]["source_id"]
    assert empty_source_id in record["topic_review_source_ids"]
    assert not any(
        candidate["relative_path"] == "empty.md"
        for candidate in record["topic_candidates"]
    )

    units_by_id = {unit.unit_id: unit for unit in first.source_units}
    for chunk in first.chunks:
        unit = units_by_id[chunk.unit_id]
        assert 0 <= chunk.char_start <= chunk.char_end <= len(unit.normalized_text)
        assert chunk.text == unit.normalized_text[chunk.char_start : chunk.char_end]

    assert any(
        chunk.relative_path == "report.pdf" and chunk.locator.get("page_label") == "1"
        for chunk in first.chunks
    )
    assert any(
        chunk.relative_path == "summary.docx"
        and ("category" in chunk.locator or chunk.locator.get("page_label") == 1)
        for chunk in first.chunks
    )
    assert any(
        chunk.relative_path == "people.xlsx"
        and chunk.locator.get("sheet_name") == "Roster"
        and chunk.locator.get("row_number") == 2
        for chunk in first.chunks
    )

    assert all(
        not Path(source["relative_path"]).is_absolute() for source in record["sources"]
    )


def test_unmappable_chunk_offsets_keep_source_and_request_review(tmp_path, monkeypatch):
    from kotaemon.indices.knowledge.evaluation import local_ingest

    root = tmp_path / "sources"
    root.mkdir()
    (root / "notes.md").write_text(
        "# Notes\n\nEvidence remains available.", encoding="utf-8"
    )

    def fail_offset_mapping(_text, _splitter):
        raise ValueError("Token splitter returned a segment outside its source unit")

    monkeypatch.setattr(local_ingest, "_split_with_offsets", fail_offset_mapping)

    summary = build_local_draft(root, tmp_path / "draft")

    assert summary.source_units
    assert not summary.chunks
    assert len(summary.quality.chunk_mapping_issues) == 1
    issue = summary.quality.chunk_mapping_issues[0]
    assert issue.relative_path == "notes.md"
    assert issue.reason == "offset_mapping_failed"
    assert issue.needs_review


def test_normalization_preserves_paragraphs_and_baseline_splits(tmp_path):
    from kotaemon.indices.knowledge.evaluation import local_ingest

    root = tmp_path / "sources"
    root.mkdir()
    paragraphs = [
        f"Paragraph {index}: "
        + ("evidence token context " * 32)
        + f"\r\nContinuation for paragraph {index}."
        for index in range(40)
    ]
    source_text = (
        "# Notes\r\n\r\n"
        "Cafe\u0301 evidence on line one.\r\n"
        "Continuation on line two.\r\n\r\n" + "\r\n\r\n".join(paragraphs)
    )
    (root / "notes.md").write_bytes(source_text.encode("utf-8"))

    first = build_local_draft(root, tmp_path / "draft-a")
    second = build_local_draft(root, tmp_path / "draft-b")

    unit = next(unit for unit in first.source_units if unit.relative_path == "notes.md")
    expected_normalized = unicodedata.normalize(
        "NFC", unit.text.replace("\r\n", "\n").replace("\r", "\n")
    )
    assert unit.normalized_text == expected_normalized
    assert "\n\n" in unit.normalized_text
    assert "\nContinuation on line two." in unit.normalized_text
    assert "Café evidence" in unit.normalized_text
    assert "Cafe\u0301 evidence" not in unit.normalized_text

    baseline_chunks = local_ingest._token_splitter()._obj.split_text(
        unit.normalized_text
    )
    chunks = [chunk for chunk in first.chunks if chunk.relative_path == "notes.md"]
    assert [chunk.text for chunk in chunks] == baseline_chunks
    assert len(chunks) > 1
    assert any(
        left.char_end > right.char_start for left, right in zip(chunks, chunks[1:])
    )
    for chunk in chunks:
        assert 0 <= chunk.char_start <= chunk.char_end <= len(unit.normalized_text)
        assert chunk.text == unit.normalized_text[chunk.char_start : chunk.char_end]
    assert [
        chunk.chunk_id for chunk in first.chunks if chunk.relative_path == "notes.md"
    ] == [
        chunk.chunk_id for chunk in second.chunks if chunk.relative_path == "notes.md"
    ]


def test_parser_fallback_is_recorded(tmp_path, monkeypatch):
    from kotaemon.base import Document
    from kotaemon.indices.knowledge.evaluation import local_ingest

    root = tmp_path / "sources"
    root.mkdir()
    (root / "fallback.docx").write_bytes(b"synthetic upload")

    def synthetic_reader(source, _path):
        diagnostic = local_ingest._reader_diagnostic(
            source,
            "DocxReader(local fallback)",
            ("UnstructuredReader(split_documents=True)", "DocxReader(local fallback)"),
            True,
            "document",
            parser_version=None,
            fallback_reason="preferred_reader_import_error",
        )
        return (
            [
                Document(
                    text="A synthetic fallback parser output.",
                    metadata={"category": "NarrativeText"},
                )
            ],
            diagnostic,
            None,
        )

    monkeypatch.setattr(local_ingest, "_load_reader_documents", synthetic_reader)
    summary = build_local_draft(root, tmp_path / "draft")

    diagnostic = summary.quality.reader_diagnostics[0]
    assert diagnostic.parser_version is None
    assert diagnostic.fallback_reason == "preferred_reader_import_error"
    assert diagnostic.extraction_granularity == "document"
    records = json.loads(summary.records_payload)
    source = records["sources"][0]
    assert source["reader_parser_version"] is None
    assert source["reader_fallback_reason"] == "preferred_reader_import_error"
    assert source["reader_extraction_granularity"] == "document"
